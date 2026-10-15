"""Arma las planillas de datos/ejemplo con personas inventadas.

Copian la forma de las reales —las columnas del export de Bombo y el desorden
de una lista armada a mano— pero ningún nombre, mail ni DNI es de alguien.
Las bases reales no se suben nunca al repositorio.

python scripts/generar_ejemplos.py
"""

from pathlib import Path

from openpyxl import Workbook

DESTINO = Path(__file__).resolve().parent.parent / "datos" / "ejemplo"

ENCABEZADO_BOMBO = [
    "Order Ticket Id", "Ticket Type", "Ticket Price", "Ticket Name", "Buyer First Name",
    "Buyer Last Name", "Buyer Email", "Buyer Phone Number", "Buyer Dni", "Order Status",
    "Last User Email", "Last User First Name", "Last User Last Name", "Last User Gender",
    "Team Lead First Name", "Team Lead Last Name",
]

ANA = ("Ana", "Ejemplo", "ana.ejemplo@example.com", 40111222, None, "Female")
DIEGO = ("Diego", "Ficticio", "diego.ficticio@example.com", 42333444, "+54 9 11 5555-0001", "Male")
PRODUCTORA = ("Bestas", "Productions")
RRPP = ("Rocío", "Vendedora")  # una RRPP inventada

# (tanda, precio, comprador, asistente o None si es el mismo, vendedor)
# comprador: nombre, apellido, mail, DNI, teléfono, género que Bombo anota en el asistente
ENTRADAS = [
    ("Early Bird", 22900, ANA, None, PRODUCTORA),
    ("Early Bird", 22900, ANA, ("Bruno", "Prueba", "bruno.prueba@example.com", "Male"), PRODUCTORA),
    ("Tier 1", 25900, ("carla", "inventada", "Carla.Inventada@Example.com ", 41222333, None, "Female"), None, PRODUCTORA),
    ("Tier 1 Combo 2 Pax", 23900, DIEGO, None, RRPP),
    ("Tier 1 Combo 2 Pax", 23900, DIEGO, ("Elena", "Muestra", None, "Mujer"), RRPP),
    ("Tier 2 VIP STANDING", 48900, ("Lucía", "Demo", "lucia.demo@example.com", 43444555, None, "Female"), None, PRODUCTORA),
]

# Una segunda fecha: Ana y Bruno vuelven, y llega alguien nuevo.
ENTRADAS_SEGUNDA = [
    ("Early Bird", 27900, ANA, None, RRPP),
    ("Tier 1", 30900, ("Bruno", "Prueba", "bruno.prueba@example.com", 44555666, None, "Male"), None, PRODUCTORA),
    ("Tier 1", 30900, ("Julia", "Nueva", "julia.nueva@example.com", 46777888, None, "Female"), None, RRPP),
]


def bombo(destino=DESTINO, entradas=ENTRADAS, archivo="bombo-ejemplo.xlsx", prefijo="ejemplo"):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "EVENTO EJEMPLO"
    hoja.append(ENCABEZADO_BOMBO)
    for i, (tanda, precio, comprador, asistente, vendedor) in enumerate(entradas, start=1):
        nombre, apellido, email, dni, telefono, genero = comprador
        if asistente:
            a_nombre, a_apellido, a_email, a_genero = asistente
        else:
            a_nombre, a_apellido, a_email, a_genero = nombre, apellido, None, genero
        hoja.append([
            f"{prefijo}{i:04d}", "Estandar", float(precio), tanda, nombre, apellido, email,
            telefono, float(dni), "approved", a_email, a_nombre, a_apellido, a_genero, *vendedor,
        ])
    ruta = destino / archivo
    libro.save(ruta)
    return ruta


def bombo_segunda(destino=DESTINO):
    return bombo(destino, ENTRADAS_SEGUNDA, "bombo-ejemplo-2.xlsx", "segunda")


def lista(destino=DESTINO):
    """Como las listas reales: dos columnas de nombre y mail, títulos sueltos y basura."""
    libro = Workbook()
    hoja = libro.active
    hoja["E1"] = "MUJERES"
    hoja["C3"] = "BASE DE DATOS"
    hoja["C5"] = "LISTA EJEMPLO"
    mujeres = [
        ("Ana Ejemplo", "ana.ejemplo@example.com"),          # también compró en Bombo
        ("Lucía Demo:", "LUCIA.DEMO@example.com"),           # dos puntos y mayúsculas
        ("Florencia Prueba", "florencia.prueba@example.com"),
        ("Florencia Prueba", "florencia.prueba@example.com"),  # repetida
        ("Martina Muestra", "martina.muestra@example"),      # mail sin dominio
    ]
    hombres = [
        ("Bruno Prueba", "bruno.prueba@example.com"),        # también asistió en Bombo
        ("tomás ficticio", "tomas.ficticio@example.com"),
    ]
    for fila, (nombre, email) in enumerate(mujeres, start=3):
        hoja.cell(fila, 5, nombre)
        hoja.cell(fila, 6, email)
    for fila, (nombre, email) in enumerate(hombres, start=3):
        hoja.cell(fila, 9, nombre)
        hoja.cell(fila, 10, email)
    hoja["N9"] = "Nombre Sin Mail"
    ruta = destino / "lista-ejemplo.xlsx"
    libro.save(ruta)
    return ruta


def clientes(destino=DESTINO):
    """Como las bases de clientes reales: varias hojas, encabezado con DNI y notas sueltas."""
    libro = Workbook()
    rrpp = libro.active
    rrpp.title = "RRPP"
    for fila in [
        ("Bruno Prueba", "bruno.prueba@example.com"),        # también en las otras planillas
        ("Gabriela Ensayo (BOMBO) 👑", "gabriela.ensayo@example.com"),  # quién la trajo, pegado al nombre
        ("Medio Ejemplo", "no comisiona"),                   # una nota donde iba el mail
    ]:
        rrpp.append(fila)

    clientes = libro.create_sheet("CLIENTES")
    clientes.append(["Nombres", "Mail", "DNI"])
    for fila in [
        ("HERNÁN DE LA MUESTRA", "hernan.muestra@example.com", 45555666.0),
        ("IGNACIO SIN MAIL", None, None),                    # en mayúsculas pero en la columna de nombres
        (None, "sin.nombre@example.com", None),
        ("Venta en puerta", "-", None),
        ("Carla Inventada", "carla.inventada@example.com // carla.otra@example.com", None),  # dos mails
    ]:
        clientes.append(fila)
    ruta = destino / "clientes-ejemplo.xlsx"
    libro.save(ruta)
    return ruta


if __name__ == "__main__":
    DESTINO.mkdir(parents=True, exist_ok=True)
    bombo()
    bombo_segunda()
    lista()
    clientes()
    print(f"Planillas de ejemplo en {DESTINO}")
