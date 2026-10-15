"""Importadores: uno por cada forma en que llegan las planillas."""

import re

from openpyxl import load_workbook

from . import normalizar as n
from .personas import anotar_incidencia, anotar_origen, guardar

# ---------------------------------------------------------------- Bombo

COLUMNAS_BOMBO = {
    "id": "Order Ticket Id",
    "tanda": "Ticket Name",
    "precio": "Ticket Price",
    "estado": "Order Status",
    "nombre": "Buyer First Name",
    "apellido": "Buyer Last Name",
    "email": "Buyer Email",
    "telefono": "Buyer Phone Number",
    "dni": "Buyer Dni",
    "asistente_email": "Last User Email",
    "asistente_nombre": "Last User First Name",
    "asistente_apellido": "Last User Last Name",
    "asistente_genero": "Last User Gender",
}
# Las que no son obligatorias: si el export no las trae, la entrada se carga igual.
COLUMNAS_BOMBO_OPCIONALES = {
    "vendedor_nombre": "Team Lead First Name",  # quien vendió: la productora o un RRPP
    "vendedor_apellido": "Team Lead Last Name",
}


def slug(texto, fecha=""):
    return re.sub(r"[^a-z0-9]+", "-", f"{texto} {fecha}".lower()).strip("-")


def importar_bombo(conn, ruta, evento, fecha):
    """Carga el export de entradas de Bombo de un evento.

    Cada entrada tiene un comprador (siempre con mail y DNI) y un asistente,
    que es a quien se le pasó la entrada. Si el asistente no tiene mail es el
    mismo comprador; si tiene otro nombre y no tiene mail, queda una incidencia.
    Si el export trae el 'Team Lead', se guarda como vendedor de la entrada.
    """
    fuente = f"Bombo · {evento}"
    evento_slug = slug(evento, fecha)
    conn.execute(
        "INSERT OR IGNORE INTO eventos (slug, nombre, fecha) VALUES (?, ?, ?)",
        (evento_slug, evento, fecha),
    )
    filas = load_workbook(ruta, read_only=True, data_only=True).active.iter_rows(values_only=True)
    encabezado = [n.texto(c) for c in next(filas)]
    faltan = [c for c in COLUMNAS_BOMBO.values() if c not in encabezado]
    if faltan:
        raise ValueError(f"No parece un export de Bombo: faltan las columnas {', '.join(faltan)}")
    indice = {
        clave: encabezado.index(columna)
        for clave, columna in {**COLUMNAS_BOMBO, **COLUMNAS_BOMBO_OPCIONALES}.items()
        if columna in encabezado
    }

    resumen = {"entradas": 0, "personas nuevas": 0, "incidencias": 0}
    for numero, fila in enumerate(filas, start=2):
        dato = {clave: fila[i] for clave, i in indice.items()}
        if not n.texto(dato["id"]):
            continue
        ubicacion = f"fila {numero}"

        comprador = n.nombre(dato["nombre"], dato["apellido"])
        mail = _mail(conn, fuente, ubicacion, dato["email"], resumen)
        documento = n.dni(dato["dni"])
        if not mail and not documento:
            anotar_incidencia(conn, fuente, ubicacion, "comprador sin mail ni DNI", comprador or "")
            resumen["incidencias"] += 1
            continue
        comprador_id, nueva = guardar(
            conn, email=mail, nombre=comprador, dni=documento,
            telefono=n.telefono(dato["telefono"]),
        )
        resumen["personas nuevas"] += nueva
        anotar_origen(conn, comprador_id, fuente, "compró")
        resumen["incidencias"] += _revisar_nombre(conn, fuente, ubicacion, comprador)

        asistente = n.nombre(dato["asistente_nombre"], dato["asistente_apellido"])
        asistente_mail = _mail(conn, fuente, f"{ubicacion}, asistente", dato["asistente_email"], resumen)
        genero = n.genero(dato["asistente_genero"])
        if asistente_mail:
            asistente_id, nueva = guardar(conn, email=asistente_mail, nombre=asistente, genero=genero)
            resumen["personas nuevas"] += nueva
        elif not asistente or asistente.lower() == (comprador or "").lower():
            asistente_id = comprador_id
            guardar(conn, email=mail, dni=documento, genero=genero)
        else:
            asistente_id = None
            anotar_incidencia(conn, fuente, ubicacion, "asistente sin mail", asistente)
            resumen["incidencias"] += 1
        if asistente_id:
            anotar_origen(conn, asistente_id, fuente, "asistió")

        vendedor = n.nombre(dato.get("vendedor_nombre"), dato.get("vendedor_apellido"))
        conn.execute(
            "INSERT INTO entradas (id, evento_slug, comprador_id, asistente_id, tanda, precio, estado, vendedor) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?) "
            "ON CONFLICT (id) DO UPDATE SET asistente_id = excluded.asistente_id, "
            "estado = excluded.estado, vendedor = excluded.vendedor",
            (n.texto(dato["id"]), evento_slug, comprador_id, asistente_id,
             n.texto(dato["tanda"]), float(dato["precio"] or 0), n.texto(dato["estado"]), vendedor),
        )
        resumen["entradas"] += 1
    conn.commit()
    return resumen


# ---------------------------------------------------------------- listas sueltas

# Los títulos de columna que dicen el género, escritos como pueden venir en una
# planilla. En la base el origen se anota siempre igual: 'columna mujeres' u 'hombres'.
TITULOS_CON_GENERO = {
    "mujeres": "mujer", "mujer": "mujer", "chicas": "mujer", "minas": "mujer",
    "hombres": "hombre", "hombre": "hombre", "chicos": "hombre", "varones": "hombre",
    "chabones": "hombre",
}
COLUMNA_DE_GENERO = {"mujer": "mujeres", "hombre": "hombres"}


def importar_lista(conn, ruta, fuente):
    """Carga una planilla armada a mano: pares de nombre y mail en cualquier lugar de la hoja.

    Lee todas las hojas del archivo. No hace falta que las columnas tengan
    encabezado: cada celda con un mail es un contacto, y su nombre es la celda
    de la izquierda. Si la columna tiene un título arriba ('MUJERES'), se anota en
    el origen y, si dice el género, se usa. Si hay una columna titulada 'DNI',
    se toma de ahí. Un nombre sin mail no se puede cargar: queda como incidencia.
    """
    hojas = load_workbook(ruta, read_only=True, data_only=True).worksheets
    resumen = {"contactos": 0, "personas nuevas": 0, "repetidos": 0, "incidencias": 0}
    vistos = set()
    for hoja in hojas:
        _importar_hoja(conn, hoja, fuente, resumen, vistos, nombrar_hoja=len(hojas) > 1)
    conn.commit()
    return resumen


def _importar_hoja(conn, hoja, fuente, resumen, vistos, nombrar_hoja):
    celdas = {
        (fila, columna): limpio
        for fila, valores in enumerate(hoja.iter_rows(values_only=True), start=1)
        for columna, valor in enumerate(valores, start=1)
        if (limpio := n.texto(valor))
    }
    con_mail = [(fila, columna) for (fila, columna), valor in celdas.items() if "@" in valor]
    if not con_mail:
        return  # una hoja sin ningún mail no es una lista de contactos
    columnas_mail = {columna for _, columna in con_mail}
    desde = min(fila for fila, _ in con_mail)  # lo de más arriba son títulos y encabezados
    columna_dni = next(
        (c for (f, c), valor in celdas.items() if f < desde and valor.lower() in ("dni", "documento")),
        None,
    )
    titulos = _titulos_de_columna(celdas)
    prefijo = f"{hoja.title}!" if nombrar_hoja else ""

    for (fila, columna), valor in sorted(celdas.items()):
        if fila < desde or (fila, columna) in titulos.values() or columna == columna_dni:
            continue
        ubicacion = prefijo + _coordenada(fila, columna)
        izquierda = celdas.get((fila, columna - 1))
        derecha = celdas.get((fila, columna + 1))

        if "@" in valor:
            encontrados = n.mails(valor)
            if not encontrados:
                anotar_incidencia(conn, fuente, ubicacion, "mail inválido", valor)
                resumen["incidencias"] += 1
                continue
            mail, *otros = encontrados
            for corregido in encontrados:
                if corregido not in valor.lower().replace(" ", ""):
                    anotar_incidencia(conn, fuente, ubicacion, "mail corregido", f"{valor} -> {corregido}")
                    resumen["incidencias"] += 1
            if mail in vistos:
                resumen["repetidos"] += 1
            vistos.add(mail)
            titulo = celdas.get(titulos.get(columna - 1))
            genero = TITULOS_CON_GENERO.get((titulo or "").lower())
            persona, nota = n.separar_nota(izquierda) if izquierda and "@" not in izquierda else (None, None)
            persona = n.nombre(persona)
            if otros:
                nota = "; ".join(filter(None, [nota, "otro mail: " + ", ".join(otros)]))
            persona_id, nueva = guardar(
                conn, email=mail, nombre=persona, notas=nota,
                dni=n.dni(celdas.get((fila, columna_dni))) if columna_dni else None,
                genero=genero, genero_origen="lista",
            )
            detalle = " · ".join(filter(None, [
                f"hoja {hoja.title}" if nombrar_hoja else None,
                f"columna {COLUMNA_DE_GENERO.get(genero, titulo)}" if titulo else None,
            ]))
            anotar_origen(conn, persona_id, fuente, detalle or "sin título")
            resumen["contactos"] += 1
            resumen["personas nuevas"] += nueva
            resumen["incidencias"] += _revisar_nombre(conn, fuente, ubicacion, persona)
        elif derecha and "@" in derecha:
            continue  # el nombre de un contacto: se cargó con su mail
        elif columna in columnas_mail or valor.isdigit():
            continue  # una nota donde iba el mail ('-', 'no comisiona'): sale con el nombre
        elif columna + 1 in columnas_mail or not valor.isupper():
            # En la columna de nombres, sin mail al lado; o un nombre suelto que no es
            # un título en mayúsculas ('BASE DE DATOS').
            detalle = f"{valor} (al lado: {derecha})" if derecha else valor
            anotar_incidencia(conn, fuente, ubicacion, "nombre sin mail", detalle)
            resumen["incidencias"] += 1


def _titulos_de_columna(celdas):
    """{columna de nombres: (fila, columna) del título} para las que tienen uno arriba.

    El título es la primera celda con texto encima de la lista que no tiene un
    mail al lado: en una de las listas reales, el que dice que son mujeres.
    """
    primer_mail = {}
    for (fila, columna), valor in sorted(celdas.items()):
        if "@" in valor:
            primer_mail.setdefault(columna, fila)
    titulos = {}
    for columna_mail, desde in primer_mail.items():
        columna_nombre = columna_mail - 1
        for fila in range(1, desde):
            if (fila, columna_nombre) in celdas and (fila, columna_mail) not in celdas:
                titulos[columna_nombre] = (fila, columna_nombre)
                break
    return titulos


def _coordenada(fila, columna):
    """(3, 5) -> 'E3', como se ve en Excel."""
    letras = ""
    while columna:
        columna, resto = divmod(columna - 1, 26)
        letras = chr(65 + resto) + letras
    return f"{letras}{fila}"


def _mail(conn, fuente, ubicacion, valor, resumen):
    """El mail normalizado; si hubo que corregirle el dominio, lo deja anotado para revisar."""
    mail, corregido = n.email_corregido(valor)
    if corregido:
        anotar_incidencia(conn, fuente, ubicacion, "mail corregido", f"{n.texto(valor)} -> {mail}")
        resumen["incidencias"] += 1
    return mail


def _revisar_nombre(conn, fuente, ubicacion, nombre):
    if n.esta_roto(nombre):
        anotar_incidencia(conn, fuente, ubicacion, "nombre con caracteres rotos", nombre)
        return 1
    return 0
